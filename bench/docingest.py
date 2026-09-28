"""Report documents (PDF, PPTX) -> Markdown per page, with every table read from a picture of it and checked.

PDF: the text layer gives the prose. Each table region (ruled tables, runs of number-heavy lines, and
embedded images, which are often tables pasted as pictures) is cropped, rendered at 200 dpi and read by a
vision model into a Markdown table. Then it is checked:
  - text-layer tables: every row's numbers must appear on one line of the page's own text in that region,
    and every number in the region must be in the transcription (catches dropped or misread digits)
  - picture-only tables: a second, independent read by the reviewer model (it doesn't see the first read);
    the two reads are compared row by row, number by number
A table is "verified" when its check passes, otherwise "flagged" with the differences, for a person to settle.
Pages with no text layer (scans) are read whole from an image and checked the same way as pictures.

PPTX: text boxes, native tables and chart data come straight from the file (exact, no reading involved);
pictures go through the same read-and-check as PDF pictures.

process() returns {"kind", "pages": [{"n", "title", "md"}], "tables": [...], "markdown"} and writes table
images to out_dir/tables/. The Markdown has <!-- page N --> markers so extracted facts can cite pages.
    uv run python bench/docingest.py report.pdf out/docs/x [model]
"""
import base64
import io
import json
import re
import sys
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import calllog
import lessons

DPI = 200
MIN_IMAGE_PT = (150, 60)  # smaller embedded images are logos / icons
MIN_PICTURE_PX = (300, 90)
READERS = 4  # table reads in flight at once (each call still waits for rate-limit room)

READ_PROMPT = """This image is cut from {where} of a valuation report. If it shows a table, transcribe it
exactly as one GitHub-flavoured Markdown table: every row and column, numbers exactly as printed (keep thousands
separators, decimals, %, currency symbols and brackets for negatives), empty cells left empty, a header row
first (repeat a merged header in each column it spans). Put the table's caption or title, if one is visible,
in title. If it isn't a table (a chart, photo, logo or diagram), set is_table false, markdown "" and describe it
in one sentence.

Rules learned from earlier reviews (apply where relevant):
{rules}"""
PAGE_PROMPT = """This is {where} of a valuation report, as an image. Transcribe it to Markdown: headings as ###,
paragraphs as text, bullet lists as lists, and every table as a GitHub-flavoured Markdown table with numbers
exactly as printed. Don't summarise or add commentary. Put "" in title if there's no heading.

Rules learned from earlier reviews (apply where relevant):
{rules}"""
FIX_PROMPT = """You transcribed the table in this image, cut from {where} of a valuation report, and checks found the
problems below. Look at the image again and correct the transcription so it matches the image exactly: every
row and column, numbers exactly as printed (thousands separators, decimals, %, currency symbols, brackets for
negatives), headings and units as printed, empty cells left empty, a header row first. Return the whole table.
{page_text}
Problems:
{problems}

Your transcription:
{markdown}

Rules learned from earlier reviews (list the IDs you apply in rules_applied):
{rules}"""
VERIFY_PROMPT = """You are the reviewer. Another model transcribed the table in this image, cut from {where} of a
valuation report, and has just corrected it for the problems below. Check the transcription against the image:
for the rows and headings involved, every number, label and unit must match the image and each value must sit
under the right column. Accept only if it does; otherwise object and list what is still wrong (the row, what the
transcription says and what the image shows).

Problems it was corrected for:
{problems}

Transcription:
{markdown}

Rules learned from earlier reviews (list the IDs you apply in rules_applied):
{rules}"""
ARBITER_PROMPT = """You are an independent arbiter. Two models read the table in this image, cut from {where} of a
valuation report, and a review loop couldn't settle it{why}. Look at the image yourself. For every row in question
(where the two reads differ, or a check below objects), write out what the image shows for that row (every value
in column order, exactly as printed) and say which read has it right. Then give each model one short piece of
feedback: what to look at again and what to correct. Other mentions of these figures in the report are below;
use one only if it clearly states the same item (same measure, same column, same date, same entity): a number
that merely coincides proves nothing.

Open points:
{problems}

The extractor's latest transcription:
{first}

The reviewer's independent read:
{second}
{page_text}{mentions}
Rules learned from earlier reviews (list the IDs you apply in rules_applied):
{rules}"""
MENTIONS_PROMPT = """A table in a valuation report ({where}) was read from its image, and the figures below are in
question: the reads of the image and the PDF's own text disagree on them, or the two reads disagree with each
other. Other parts of the report contain the same numbers. For each passage, decide whether it states one of
these figures for the same item: the same measure, the same column (low / mid / high, a year, a scenario), the
same date and the same entity. If it does, say whether it confirms what the reads of the image show, or
contradicts them. A number that merely coincides, or the same measure for another period, scenario or entity,
is "not the same item". When unsure, say "not the same item".

Table: {title}
{table}
Figures in question:
{disputed}

Passages:
{passages}"""
FEEDBACK_NOTE = """
An independent arbiter looked at this image too and noted the points below. Check each against the image
yourself; transcribe what the image shows, not what the note says.
{feedback}
"""
MAX_ROUNDS = 3
VISUAL_AFTER = 2  # rounds the extractor gets against the page's text layer before two agreeing reads of the image
                  # outrank it (the text layer is usually right, but PDFs do glue or split characters)
LOOP_VERSION = 2  # bump when the loop changes: tables it left for a person go round once more
                  # (2: agreeing reads of the image outrank the text layer; the arbiter; the report's other mentions)
_S = {"type": "string"}
_READ = {"type": "json_schema", "name": "table_read", "strict": True, "schema": {
    "type": "object", "additionalProperties": False, "required": ["is_table", "title", "markdown", "description"],
    "properties": {"is_table": {"type": "boolean"}, "title": _S, "markdown": _S, "description": _S}}}
_PAGE = {"type": "json_schema", "name": "page_read", "strict": True, "schema": {
    "type": "object", "additionalProperties": False, "required": ["title", "markdown"],
    "properties": {"title": _S, "markdown": _S}}}
_IDS = {"type": "array", "items": _S}
_FIX = {"type": "json_schema", "name": "table_fix", "strict": True, "schema": {
    "type": "object", "additionalProperties": False, "required": ["markdown", "changes", "rules_applied"],
    "properties": {"markdown": _S, "changes": _S, "rules_applied": _IDS}}}
_PROBLEM = {"type": "object", "additionalProperties": False, "required": ["row", "issue"], "properties": {"row": _S, "issue": _S}}
_ARB_ROW = {"type": "object", "additionalProperties": False, "required": ["row", "image_shows", "right", "note"],
            "properties": {"row": _S, "image_shows": _S, "note": _S,
                           "right": {"type": "string", "enum": ["extractor", "reviewer", "both", "neither"]}}}
_ARBITER = {"type": "json_schema", "name": "table_arbiter", "strict": True, "schema": {
    "type": "object", "additionalProperties": False,
    "required": ["rows", "to_extractor", "to_reviewer", "rules_applied"],
    "properties": {"rows": {"type": "array", "items": _ARB_ROW}, "to_extractor": _S, "to_reviewer": _S,
                   "rules_applied": _IDS}}}
_MENTION = {"type": "object", "additionalProperties": False, "required": ["passage", "verdict", "figure", "why"],
            "properties": {"passage": {"type": "integer"}, "figure": _S, "why": _S,
                           "verdict": {"type": "string", "enum": ["confirms", "contradicts", "not the same item"]}}}
_MENTIONS = {"type": "json_schema", "name": "table_mentions", "strict": True, "schema": {
    "type": "object", "additionalProperties": False, "required": ["passages"],
    "properties": {"passages": {"type": "array", "items": _MENTION}}}}
_VERIFY = {"type": "json_schema", "name": "table_verify", "strict": True, "schema": {
    "type": "object", "additionalProperties": False, "required": ["verdict", "problems", "rules_applied"],
    "properties": {"verdict": {"type": "string", "enum": ["accept", "object"]},
                   "problems": {"type": "array", "items": _PROBLEM}, "rules_applied": _IDS}}}


# ---- numbers ------------------------------------------------------------------------------------------------

NUM = re.compile(r"\(?[-−–]?\s?(?:[A-Z]{0,2}\$|€|£)?\s?\d[\d,]*(?:\.\d+)?\s?%?\)?")


def numbers(text: str) -> list[str]:
    """Numbers as printed, normalised: '(1,234.5)' -> '-1234.5', '7.25%' -> '7.25%', 'A$850.0m' -> '850.0'."""
    out = []
    for m in NUM.finditer(text or ""):
        s = m.group(0).strip()
        digits = re.sub(r"[^\d.]", "", s).rstrip(".")
        if not digits or not re.search(r"\d", digits):
            continue
        neg = (s.startswith("(") and s.endswith(")")) or bool(re.match(r"\(?[-−–]", s))
        out.append(("-" if neg else "") + digits + ("%" if "%" in s else ""))
    return out


def md_rows(md: str) -> list[list[str]]:
    """Cells of each body row of the Markdown table(s) in md (separator rows skipped)."""
    rows = []
    for line in (md or "").splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def _row_key(cells: list[str]) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (cells[0] if cells else "").lower()).strip()


def compare_reads(a: str, b: str) -> dict:
    """Two independent transcriptions of one table, compared row by row (rows matched by first cell)."""
    ra, rb = md_rows(a), md_rows(b)
    diffs = []
    kb = {}
    for cells in rb:
        kb.setdefault(_row_key(cells), []).append(cells)
    for cells in ra:
        others = kb.get(_row_key(cells))
        other = others.pop(0) if others else None
        na = numbers(" ".join(cells[1:]))
        nb = numbers(" ".join(other[1:])) if other else None
        if nb is None:
            if na:
                diffs.append({"row": cells[0], "first": " | ".join(cells[1:]), "second": "(row not found)"})
        elif Counter(na) != Counter(nb):
            diffs.append({"row": cells[0], "first": " | ".join(cells[1:]), "second": " | ".join(other[1:])})
    for rest in kb.values():
        for cells in rest:
            if numbers(" ".join(cells[1:])):
                diffs.append({"row": cells[0], "first": "(row not found)", "second": " | ".join(cells[1:])})
    n = sum(len(numbers(" ".join(c[1:]))) for c in ra)
    return {"method": "second read", "numbers": n, "rows": len(ra), "differences": diffs, "ok": not diffs and n > 0}


def _in_order(seq: list, line: list) -> bool:
    it = iter(line)
    return all(x in it for x in seq)


def _words(text: str) -> Counter:
    return Counter(w for w in re.findall(r"[a-z]+", (text or "").lower()))


def _sq(s: str) -> str:
    return re.sub(r"\s+", "", s or "").lower()


_NOT_BEFORE, _NOT_AFTER = set("0123456789.,(-−–"), set("0123456789%)")


def _find(sq: str, want: str, start: int = 0) -> int:
    """want in sq from start, not as part of a longer number ("2" inside "12" or "2.5", "5" inside "5%", "12"
    inside "(12)"); -1 if not there."""
    j = sq.find(want, start)
    while j >= 0:
        e = j + len(want)
        before, after = sq[j - 1] if j else "", sq[e] if e < len(sq) else ""
        if before not in _NOT_BEFORE and after not in _NOT_AFTER and not (after in ".," and sq[e + 1:e + 2].isdigit()):
            return j
        j = sq.find(want, j + 1)
    return -1


def _take(left: list[list[str]], cells: list[str]) -> bool:
    """Find a row's cells, spacing ignored, in order on one line of the page; mark those characters used (so no
    other row can claim them) and say whether they were found. The characters must be the same: only spaces may
    differ ("3. 75%" is 3.75%, "Hi gh" is High). Cells are looked for one after another, other text allowed
    between them; failing that, run together ("222" on the page for 2 | 2 | 2)."""
    want = [w for w in map(_sq, cells) if w]
    if not want:
        return False
    for chars in left:
        idx = [i for i, ch in enumerate(chars) if not ch.isspace()]
        sq = "".join(chars[i] for i in idx).lower()
        spans, pos = [], 0
        for w in want:
            j = _find(sq, w, pos)
            if j < 0:
                break
            spans.append((j, j + len(w)))
            pos = j + len(w)
        else:
            for a, b in spans:
                for k in range(a, b):
                    chars[idx[k]] = "\x00"
            return True
        j = _find(sq, "".join(want))
        if j >= 0:
            for k in range(j, j + len("".join(want))):
                chars[idx[k]] = "\x00"
            return True
    return False


def _letters(text: str) -> str:
    return re.sub(r"[^a-z]", "", (text or "").lower())


def check_text_layer(md: str, lines: list[str], title: str | None = None) -> dict:
    """Transcription vs the PDF's own text in the table region. lines = the region's text, one entry per line.
    Numbers must match exactly; words too, so a dropped unit ("$m" for "A$m") or a misread label is caught.
    Only spacing is forgiven, because PDFs often get it wrong: letter-spaced words ("Hi gh"), numbers split at the
    decimal point ("3. 75%"), or run together across columns ("222" for 2 | 2 | 2) still match when the characters
    are the same and in the same order on one line. The table's title may account for words of its caption."""
    left = [list(ln) for ln in lines]
    src_lines = [numbers(ln) for ln in lines]
    got, total, bad_rows = Counter(), 0, []
    for cells in md_rows(md):
        row, every = numbers(" ".join(cells[1:])), numbers(" ".join(cells))
        total += len(every)
        if every and _take(left, [c for c in cells if c]):  # the whole row, label included (or no label at all)
            continue
        if row and (_take(left, [c for c in cells[1:] if c]) or _take(left, [c for c in cells[1:] if numbers(c)])):
            got += Counter(numbers(cells[0]))  # the label's own numbers are matched one by one below
            continue
        got += Counter(every)
        # the row's numbers, in order, on one source line (so a swapped column or a row's figures under
        # another row's label is caught, not just a misread digit)
        if row and not any(_in_order(row, line) for line in src_lines):
            bad_rows.append({"row": cells[0], "numbers": row})
    src = Counter(n for chars in left for n in numbers("".join(chars)))
    not_in_source = list((got - src).elements())
    not_transcribed = list((src - got).elements())
    src_words, got_words = _words(" ".join(lines)), _words(md)
    src_letters = "|".join(_letters(ln) for ln in lines)
    got_letters = "|".join(_letters(" ".join(c)) for c in md_rows(md))
    words_not_in_source = sorted(w for w in set(got_words - src_words) if w not in src_letters)
    words_not_transcribed = sorted(w for w in set(src_words - got_words)
                                   if w not in got_letters and w not in _letters(title))
    return {"method": "text layer", "numbers": total, "rows_not_on_one_line": bad_rows,
            "not_in_source": not_in_source, "not_transcribed": not_transcribed,
            "words_not_in_source": words_not_in_source, "words_not_transcribed": words_not_transcribed,
            "ok": not (bad_rows or not_in_source or not_transcribed or words_not_in_source or words_not_transcribed)
                  and total > 0}


# ---- model calls --------------------------------------------------------------------------------------------

class Reader:
    """Vision calls through llm.create (rate-limited); usage goes to on_usage(model, usage, purpose)."""

    def __init__(self, model: str, reviewer_model: str, on_usage=None, arbiter_model: str | None = None):
        from llm import client
        self.llm, self.model, self.reviewer_model = client(interactive=False), model, reviewer_model
        self.arbiter_model = arbiter_model
        self.on_usage = on_usage or (lambda *a: None)

    def _call(self, model, prompt, png: bytes | None, schema, purpose) -> dict:
        from llm import create
        content = [{"type": "input_text", "text": prompt}]
        if png is not None:
            content.append({"type": "input_image", "image_url": f"data:image/png;base64,{base64.b64encode(png).decode()}",
                            "detail": "high"})
        r = create(self.llm, model, text={"format": schema}, max_output_tokens=6000, purpose=purpose,
                   input=[{"role": "user", "content": content}])
        if r.usage:
            self.on_usage(model, r.usage, purpose)
        return json.loads(r.output_text)

    def table(self, png: bytes, where: str, second: bool = False, feedback: str | None = None) -> dict:
        note = FEEDBACK_NOTE.format(feedback=feedback) if feedback else ""
        return self._call(self.reviewer_model if second else self.model,
                          READ_PROMPT.format(where=where, rules=lessons.rules_text("tables")) + note, png, _READ,
                          "doc-review" if second else "doc-table")

    def page(self, png: bytes, where: str, second: bool = False, feedback: str | None = None) -> dict:
        note = FEEDBACK_NOTE.format(feedback=feedback) if feedback else ""
        return self._call(self.reviewer_model if second else self.model,
                          PAGE_PROMPT.format(where=where, rules=lessons.rules_text("tables")) + note, png, _PAGE,
                          "doc-review" if second else "doc-page")

    def arbitrate(self, png: bytes, where: str, why: str, problems: list[str], first: str, second: str,
                  page_lines: list[str] | None, mentions: list[dict]) -> dict:
        page_text = ("\nThe PDF's own text inside the table (exact characters, but its spacing and reading order may "
                     "be off):\n" + "\n".join(page_lines) + "\n") if page_lines else ""
        found = ("\nOther mentions in the report:\n" + "\n".join(f"- page {m['page']} ({m['kind']}): {m['text']}"
                                                                  for m in mentions) + "\n") if mentions else ""
        return self._call(self.arbiter_model, ARBITER_PROMPT.format(
            where=where, why=why, problems="\n".join(f"- {p}" for p in problems) or "- (none listed)", first=first,
            second=second or "(none)", page_text=page_text, mentions=found, rules=lessons.rules_text("tables")),
            png, _ARBITER, "doc-arbiter")

    def mentions(self, where: str, title: str, table: str, disputed: list[str], passages: list[dict]) -> dict:
        return self._call(self.reviewer_model, MENTIONS_PROMPT.format(
            where=where, title=title or "(no title)", table=table, disputed="\n".join(f"- {d}" for d in disputed),
            passages="\n".join(f"{i}. page {m['page']} ({m['kind']}): {m['text']}" for i, m in enumerate(passages, 1))),
            None, _MENTIONS, "doc-mentions")

    def fix(self, png: bytes, where: str, problems: list[str], markdown: str, page_lines: list[str] | None) -> dict:
        page_text = ("\nThe page's own text inside the table (exact characters; the reading order may differ from the "
                     "layout):\n" + "\n".join(page_lines) + "\n") if page_lines else ""
        return self._call(self.model, FIX_PROMPT.format(where=where, page_text=page_text, problems="\n".join(
            f"- {p}" for p in problems), markdown=markdown, rules=lessons.rules_text("tables")), png, _FIX, "doc-fix")

    def verify(self, png: bytes, where: str, problems: list[str], markdown: str, feedback: str | None = None) -> dict:
        note = FEEDBACK_NOTE.format(feedback=feedback) if feedback else ""
        return self._call(self.reviewer_model, VERIFY_PROMPT.format(where=where, problems="\n".join(
            f"- {p}" for p in problems), markdown=markdown, rules=lessons.rules_text("tables")) + note, png, _VERIFY,
            "doc-verify")


def _png(img) -> bytes:
    buf = io.BytesIO()
    img.convert("RGB").save(buf, format="PNG")
    return buf.getvalue()


def read_and_check(reader: Reader, t: dict, png: bytes, text_lines: list[str] | None) -> dict:
    """Fill in t (a table record) from the reads and checks."""
    whole_page = t["source"] == "page image"
    first = reader.page(png, t["where"]) if whole_page else reader.table(png, t["where"])
    if not whole_page and not first["is_table"]:
        return {**t, "is_table": False, "markdown": "", "description": first["description"], "status": "figure",
                "check": None}
    check = check_text_layer(first["markdown"], text_lines, first.get("title")) if text_lines else None
    if check and check["ok"]:  # the PDF's own text confirms every number and word: no second read needed
        return {**t, "is_table": True, "title": first.get("title") or None, "markdown": first["markdown"],
                "check": check, "status": "verified"}
    second = reader.page(png, t["where"], True) if whole_page else reader.table(png, t["where"], True)
    agree = compare_reads(first["markdown"], second["markdown"])
    ok = check["ok"] if check else agree["ok"]
    return {**t, "is_table": True, "title": first.get("title") or None, "markdown": first["markdown"],
            "second_markdown": second["markdown"], "check": check, "second_read": agree,
            "status": "verified" if ok else "flagged"}


# ---- the review and remediation loop for flagged tables ------------------------------------------------------
# The first reader fixes its transcription against the problems found; code re-checks it against the page's text
# layer where there is one (that check is the ground truth for characters, spacing aside); the reviewer model then
# checks the fix against the image. From round VISUAL_AFTER, two reads of the image that agree outrank a text layer
# that still objects, unless another part of the report states something different for the same item. What that
# doesn't settle goes to the arbiter, whose feedback both readers get for one more try. The rest stays flagged for
# a person, with every round shown.

def latest(t: dict) -> str:
    """The table's text as it stands: settled, else the loop's latest correction, else the first read."""
    return t.get("final_markdown") or (t.get("resolution") or {}).get("markdown") or t.get("markdown") or ""


def problems(t: dict) -> list[str]:
    """A flagged table's problems, in words for the fixer (from the latest check, if the loop has run)."""
    r = t.get("resolution") or {}
    c, s, out = t.get("final_check") or r.get("check") or t.get("check"), r.get("compare") or t.get("second_read"), []
    if c and c.get("method") == "text layer" and not c.get("ok"):
        if c.get("not_in_source"):
            out.append("numbers in your transcription that aren't on the page: " + ", ".join(c["not_in_source"]))
        if c.get("not_transcribed"):
            out.append("numbers on the page missing from your transcription: " + ", ".join(c["not_transcribed"]))
        if c.get("rows_not_on_one_line"):
            out.append("rows whose numbers don't appear, in this order, on one line of the page: "
                       + "; ".join(r["row"] for r in c["rows_not_on_one_line"]))
        if c.get("words_not_in_source"):
            out.append("words in your transcription that aren't on the page: " + ", ".join(c["words_not_in_source"]))
        if c.get("words_not_transcribed"):
            out.append("words on the page missing from your transcription (check headings and units): "
                       + ", ".join(c["words_not_transcribed"]))
        if not c.get("numbers"):
            out.append("no numbers were transcribed")
    elif s and not s.get("ok"):
        out.append("an independent read of the image differs on these rows; re-read them cell by cell: "
                   + "; ".join(d["row"] for d in s["differences"]))
    return out


def _value(n: str):
    try:
        return round(float(n.rstrip("%")), 6), n.endswith("%")
    except ValueError:
        return None


def _telling(n: str) -> bool:
    """A figure worth looking for elsewhere: not a small whole number or a year, which turn up everywhere."""
    v = _value(n)
    return bool(v) and (v[1] or "." in n or not (abs(v[0]) < 10 or 1900 <= v[0] <= 2100))


def find_mentions(doc: dict | None, t: dict, figures: list[str], limit: int = 12) -> list[dict]:
    """Other places in the report that state any of these figures, compared as values (7.7% finds 7.70%): sentences
    of the running text, the table's own page first, and rows of the other tables. Code only finds them; a model
    judges whether they are about the same item."""
    want = {_value(n) for n in figures if _telling(n)}
    if not doc or not want:
        return []
    hit = lambda text: bool(want & {_value(n) for n in numbers(text)})
    out = []
    for p in doc.get("pages") or []:
        text = " ".join(str(b[2]) for b in p["blocks"] if len(b) > 2 and b[1] in ("text", "heading"))
        for sent in re.split(r"(?<=[.;:!?])\s+(?=[A-Z(•–-])", text):
            if hit(sent):
                out.append({"page": p["n"], "kind": "text", "text": sent.strip()[:400]})
    for t2 in doc.get("tables") or []:
        if t2["id"] == t["id"] or t2.get("status") in ("figure", "error"):
            continue
        rows = md_rows(latest(t2))
        for cells in rows[1:]:
            if hit(" ".join(cells)):
                out.append({"page": t2.get("page"), "kind": f"table “{t2.get('title') or t2['id']}”",
                            "text": f"columns {' | '.join(rows[0])}; row {' | '.join(cells)}"[:400]})
    seen, res = set(), []
    for m in sorted(out, key=lambda m: (m["page"] != t.get("page"), m["kind"] != "text")):
        if (m["page"], m["text"]) not in seen:
            seen.add((m["page"], m["text"]))
            res.append(m)
    return res[:limit]


def _image_over_text(reader: Reader, doc: dict | None, t: dict, md: str, chk: dict) -> dict:
    """The text layer objects but two reads of the image agree: what it objects to is recorded as overruled, and the
    report's other mentions of those figures are checked (a model judges whether each is the same item)."""
    out = {"overruled": problems({"check": chk}), "searched": 0, "confirms": [], "contradicts": []}
    rows = chk.get("rows_not_on_one_line") or []
    found = find_mentions(doc, t, chk["not_in_source"] + chk["not_transcribed"] + [n for r in rows for n in r["numbers"]])
    out["searched"] = len(found)
    if not found:
        return out
    disputed = ([f"as read from the image: {', '.join(chk['not_in_source'])}"] if chk["not_in_source"] else []) + (
        [f"what the PDF's own text has instead: {', '.join(chk['not_transcribed'])}"] if chk["not_transcribed"] else []) + [
        f"row {r['row']}: {', '.join(r['numbers'])}" for r in rows]
    try:
        got = reader.mentions(t["where"], t.get("title"), md, disputed, found)
    except Exception as e:  # the other mentions are extra evidence: without them the two reads still stand
        out["error"] = f"{type(e).__name__}: {e}"
        return out
    for x in got["passages"]:
        if 1 <= x["passage"] <= len(found) and x["verdict"] in ("confirms", "contradicts"):
            out[x["verdict"]].append({**found[x["passage"] - 1], "figure": x["figure"], "why": x["why"]})
    return out


def _contradicted(v: dict) -> list[str]:
    return [f"page {m['page']} of the report states a different figure for the same item ({m['figure']}): "
            f"“{m['text'][:200]}”. {m['why']}" for m in v["contradicts"]]


def resolve_table(reader: Reader, t: dict, png: bytes, rounds: int = MAX_ROUNDS, doc: dict | None = None) -> dict | None:
    """Run the loop on one flagged table; updates t (status "resolved" if the agents settle it) and returns the
    episode for the lessons. A table that went round before continues from its latest correction; earlier rounds
    are kept and counted. doc (the whole report) lets the loop look for the figures elsewhere in it."""
    prev = t.get("resolution") or {}
    thread = list(prev.get("rounds") or [])
    k0 = sum(isinstance(x.get("round"), int) for x in thread)
    issues = (None if prev.get("resolved") else prev.get("open")) or problems(t)
    if not issues:
        return None
    md, lines, title = latest(t), t.get("text_lines"), t.get("title")
    second = prev.get("second_markdown") or t.get("second_markdown")
    chk, basis, n0 = None, None, len(thread)
    for k in range(k0 + 1, k0 + rounds + 1):
        fix = reader.fix(png, t["where"], issues, md, lines)
        lessons.applied(fix["rules_applied"])
        rec = {"round": k, "problems": issues, "fixer": reader.model, "changes": fix["changes"]}
        thread.append(rec)
        md = fix["markdown"]
        chk = check_text_layer(md, lines, title) if lines else None
        agree = compare_reads(md, second) if second else None
        if agree is not None:
            rec["agrees_with_second_read"] = agree["ok"]
        if chk is not None:
            rec["check_ok"] = chk["ok"]
            if not chk["ok"]:
                if k >= VISUAL_AFTER and agree and agree["ok"]:  # two reads of the image agree: they outrank it
                    v = rec["image_over_text"] = _image_over_text(reader, doc, t, md, chk)
                    if not v["contradicts"]:
                        basis = "image"
                        break
                    issues = _contradicted(v)
                    rec["verdict"] = {"by": "report", "verdict": "object", "problems": issues}
                    continue
                issues = problems({"check": chk})  # the page's text says it's still wrong: back to the fixer
                rec["verdict"] = {"by": "code", "verdict": "object", "problems": issues}
                continue
        ver = reader.verify(png, t["where"], issues, md)
        lessons.applied(ver["rules_applied"])
        rec["verdict"] = {"by": reader.reviewer_model, "verdict": ver["verdict"],
                          "problems": [f"{p['row']}: {p['issue']}" for p in ver["problems"]]}
        if ver["verdict"] == "accept":
            basis = "text layer" if chk is not None else "reviewer"
            break
        issues = rec["verdict"]["problems"] or ["the reviewer objected without saying where; re-read every row"]
    if basis is None and reader.arbiter_model:
        basis, md, second, chk, issues = _arbiter(reader, doc, t, png, md, second, chk, issues, thread, len(thread) - n0)
    t["resolution"] = {"resolved": basis is not None, "basis": basis, "rounds": thread, "markdown": md, "check": chk,
                       "second_markdown": second if second != t.get("second_markdown") else None,
                       "compare": compare_reads(md, second) if second else None, "open": None if basis else issues}
    if basis:
        t.update(status="resolved", final_markdown=md, final_check=chk)
    return {"source": t["source"], "text_layer": bool(lines), "rounds": thread[n0:],
            "outcome": f"resolved ({basis})" if basis else "escalated"}


def _arbiter(reader: Reader, doc, t: dict, png: bytes, md: str, second: str | None, chk: dict | None,
             issues: list[str], thread: list, n: int):
    """The loop didn't settle it: the arbiter looks at the image, both reads, the text layer and the other mentions
    of the figures, and gives each reader feedback. The extractor corrects, the reviewer reads the image again; if
    the two now agree and the text layer agrees too (or is outranked, with nothing in the report against it),
    it's settled. Returns (basis or None, markdown, reviewer's read, check, open issues)."""
    lines, whole = t.get("text_lines"), t["source"] == "page image"
    diffs = compare_reads(md, second)["differences"] if second else []
    figures = [x for d in diffs for x in numbers(d["first"] + " " + d["second"])] + (
        (chk["not_in_source"] + chk["not_transcribed"]) if chk and not chk["ok"] else [])
    found = find_mentions(doc, t, figures)
    why = f" in {n} round{'s' if n != 1 else ''}"
    arb = reader.arbitrate(png, t["where"], why, issues, md, second, lines, found)
    lessons.applied(arb["rules_applied"])
    rows = [f"row {r['row']}: the image shows {r['image_shows']}" + (f" ({r['note']})" if r["note"] else "")
            for r in arb["rows"]]
    rec = {"round": "arbiter", "by": reader.arbiter_model, "rows": arb["rows"], "to_extractor": arb["to_extractor"],
           "to_reviewer": arb["to_reviewer"], "mentions": len(found), "problems": issues}
    thread.append(rec)
    fix = reader.fix(png, t["where"], [f"an independent arbiter looked at the image: {x}"
                                       for x in rows + [arb["to_extractor"]] if x], md, lines)
    lessons.applied(fix["rules_applied"])
    md = fix["markdown"]
    rec.update(fixer=reader.model, changes=fix["changes"])
    feedback = "\n".join(f"- {x}" for x in rows + [arb["to_reviewer"]] if x)
    second = (reader.page if whole else reader.table)(png, t["where"], True, feedback)["markdown"]
    agree = compare_reads(md, second)
    rec.update(reviewer=reader.reviewer_model, agrees_with_second_read=agree["ok"])
    chk = check_text_layer(md, lines, t.get("title")) if lines else None
    if chk is not None:
        rec["check_ok"] = chk["ok"]
    if not agree["ok"]:
        return None, md, second, chk, ["after the arbiter's feedback the two reads of the image still differ on: "
                                       + "; ".join(d["row"] for d in agree["differences"])]
    if chk is None or chk["ok"]:
        return "arbiter", md, second, chk, []
    v = rec["image_over_text"] = _image_over_text(reader, doc, t, md, chk)
    if v["contradicts"]:
        return None, md, second, chk, _contradicted(v)
    return "arbiter", md, second, chk, []


def resolve_tables(doc: dict, out_dir: str | Path, model: str, reviewer_model: str, progress=None, on_usage=None,
                   ids: set | None = None, arbiter_model: str | None = None) -> dict:
    """The loop on every flagged table of a read document (or those in ids). Returns {"resolved", "escalated",
    "episodes"}."""
    progress = progress or (lambda frac, msg: None)
    todo = [t for t in doc["tables"] if t.get("status") == "flagged" and t.get("png") and (ids is None or t["id"] in ids)]
    if not todo:
        return {"resolved": 0, "escalated": 0, "episodes": []}
    reader = Reader(model, reviewer_model, on_usage, arbiter_model)
    done, episodes = 0, []

    def one(t):
        try:
            return resolve_table(reader, t, (Path(out_dir) / t["png"]).read_bytes(), doc=doc)
        except Exception as e:  # the table stays flagged for a person
            t["resolution"] = {**(t.get("resolution") or {}), "resolved": False,
                               "open": [f"the loop failed: {type(e).__name__}: {e}"]}
            t["resolution"].setdefault("rounds", [])
            return None

    with ThreadPoolExecutor(READERS) as pool:
        for ep in pool.map(calllog.carry(one), todo):
            done += 1
            progress(done / len(todo), f"Table review loop: {done} of {len(todo)} flagged tables")
            if ep:
                episodes.append(ep)
    doc["markdown"] = render(doc)
    n = sum(t.get("status") == "resolved" for t in todo)
    return {"resolved": n, "escalated": len(todo) - n, "episodes": episodes}


def recheck(doc: dict) -> int:
    """Check every flagged text-layer table again with today's check (e.g. after it learned to forgive spacing):
    those it now passes are settled. Returns how many."""
    n = 0
    for t in doc["tables"]:
        if t.get("status") != "flagged" or not t.get("text_lines"):
            continue
        r = t.get("resolution") or {}
        if r.get("markdown"):
            chk = check_text_layer(r["markdown"], t["text_lines"], t.get("title"))
            if chk["ok"]:
                r.update(resolved=True, basis="text layer", check=chk, open=None)
                t.update(status="resolved", final_markdown=r["markdown"], final_check=chk)
                n += 1
                continue
        chk = check_text_layer(t.get("markdown") or "", t["text_lines"], t.get("title"))
        if chk["ok"]:
            t.update(status="verified", check=chk)
            n += 1
    if n:
        doc["markdown"] = render(doc)
    return n


# ---- PDF ----------------------------------------------------------------------------------------------------

def _lines(words: list[dict], tol: float = 2.5) -> list[dict]:
    """Group words into lines: [{"top", "bottom", "x0", "x1", "text", "size"}], top to bottom."""
    lines = []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1]["top"] - w["top"]) <= tol:
            ln = lines[-1]
            ln["words"].append(w)
            ln["bottom"] = max(ln["bottom"], w["bottom"])
        else:
            lines.append({"top": w["top"], "bottom": w["bottom"], "words": [w]})
    for ln in lines:
        ws = sorted(ln["words"], key=lambda w: w["x0"])
        ln.update(text=" ".join(w["text"] for w in ws), x0=ws[0]["x0"], x1=ws[-1]["x1"],
                  size=sum(w.get("size", 0) for w in ws) / len(ws),
                  bold=all(w.get("bold") for w in ws))
        del ln["words"]
    return lines


def _inside(ln: dict, box) -> bool:
    x0, top, x1, bottom = box
    mid = (ln["top"] + ln["bottom"]) / 2
    return top - 1 <= mid <= bottom + 1 and ln["x1"] >= x0 and ln["x0"] <= x1


def _number_runs(lines: list[dict], taken: list) -> list[tuple]:
    """Runs of 3+ lines with 2+ numbers each (borderless tables), plus a header line just above."""
    boxes, run = [], []

    def flush():
        if len(run) >= 3:
            i0 = lines.index(run[0])
            head = [lines[i0 - 1]] if i0 > 0 and run[0]["top"] - lines[i0 - 1]["bottom"] < 14 else []
            rs = head + run
            boxes.append((min(r["x0"] for r in rs) - 4, rs[0]["top"] - 3, max(r["x1"] for r in rs) + 4, rs[-1]["bottom"] + 3))

    for ln in lines:
        words = ln["text"].split()
        n = sum(1 for w in words if numbers(w))
        # tabular: mostly numbers (a row label plus figures), not prose that happens to quote two amounts
        numeric = n >= 2 and n >= 0.5 * len(words) and not any(_inside(ln, b) for b in taken)
        if numeric and (not run or ln["top"] - run[-1]["bottom"] < 16):
            run.append(ln)
        else:
            flush()
            run = [ln] if numeric else []
    flush()
    return boxes


# ---- reading order: columns, sidebars, text beside a chart ------------------------------------------------------
# Words on one baseline across the whole page would run two columns together line by line. The page is cut
# recursively (XY cut) into blocks read top-left to bottom-right, column by column: first down a tall clear
# gutter with running text on both sides (or a table / picture on one side), else across the widest horizontal
# gap. Lines of a label-value list or a table have too few words to count as running text, so they're never
# split into their columns.
COL_GAP = 9.0        # pt: the narrowest gutter between two columns
_BOLD = re.compile(r"bold|black|heavy", re.I)
PROSE_WORDS = 3.0    # average words per line for a side of a gutter to count as running text


def _respace(words: list[dict]) -> list[dict]:
    """Letter-spaced text (glyphs set one at a time with extra tracking, as some exporters write tables) comes out
    of the text layer as one-letter words: "N e t f i n a n c i a l d e b t 5 , 2 2 3 . 0". Per line, if most
    words are single characters, rebuild real words: letters join unless the gap is clearly wider than the line's
    usual letter gap (a word space adds about a quarter of the type size)."""
    out, lines = [], []
    for w in sorted(words, key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(lines[-1][0]["top"] - w["top"]) <= 2.5:
            lines[-1].append(w)
        else:
            lines.append([w])
    for ln in lines:
        ln.sort(key=lambda w: w["x0"])
        if len(ln) < 4 or sum(len(w["text"]) == 1 for w in ln) < 0.6 * len(ln):
            out.extend(ln)
            continue
        gaps = [b["x0"] - a["x1"] for a, b in zip(ln, ln[1:])]
        letter = sorted(gaps)[len(gaps) // 2]
        size = sorted(w.get("size", 10) for w in ln)[len(ln) // 2]
        cut = letter + max(0.25 * size, 1.0)
        run = [ln[0]]
        for w, g in zip(ln[1:], gaps):
            if g <= cut:
                run.append(w)
                continue
            out.append(_joined(run))
            run = [w]
        out.append(_joined(run))
    return out


def _joined(run: list[dict]) -> dict:
    if len(run) == 1:
        return run[0]
    w = dict(run[0])
    w.update(text="".join(x["text"] for x in run), x0=min(x["x0"] for x in run), x1=max(x["x1"] for x in run),
             top=min(x["top"] for x in run), bottom=max(x["bottom"] for x in run),
             size=sum(x.get("size", 0) for x in run) / len(run), bold=all(x.get("bold") for x in run))
    if "doctop" in w:
        w["doctop"] = min(x["doctop"] for x in run)
    w["width"], w["height"] = w["x1"] - w["x0"], w["bottom"] - w["top"]
    return w


def _merged(spans: list[tuple[float, float]]) -> list[tuple[float, float]]:
    out = []
    for a, b in sorted(spans):
        if out and a <= out[-1][1]:
            out[-1] = (out[-1][0], max(out[-1][1], b))
        else:
            out.append((a, b))
    return out


def _running_text(items: list[dict]) -> bool:
    words = [w for w in items if "box" not in w]
    if not words:
        return any("box" in w for w in items)  # a table or picture on its own
    lines = _lines(words)
    return len(lines) >= 2 and len(words) / len(lines) >= PROSE_WORDS


def _xy_cut(items: list[dict], lh: float, depth: int = 0) -> list[list[dict]]:
    """Items (words, and table / picture boxes marked "box") -> blocks in reading order."""
    if len(items) <= 1 or depth > 24:
        return [items]
    cols = _merged([(w["x0"], w["x1"]) for w in items])
    best = None
    for (_, a), (b, _) in zip(cols, cols[1:]):
        if b - a < COL_GAP:
            continue
        left, right = [w for w in items if w["x1"] <= a], [w for w in items if w["x0"] >= b]
        side_by_side = min(max(w["bottom"] for w in left), max(w["bottom"] for w in right)) - \
            max(min(w["top"] for w in left), min(w["top"] for w in right))
        if side_by_side >= 2 * lh and _running_text(left) and _running_text(right) \
                and any("box" not in w for w in left + right):
            if best is None or b - a > best[1] - best[0]:
                best = (a, b, left, right)
    if best:
        return _xy_cut(best[2], lh, depth + 1) + _xy_cut(best[3], lh, depth + 1)
    rows = _merged([(w["top"], w["bottom"]) for w in items])
    gaps = [(b - a, a, b) for (_, a), (b, _) in zip(rows, rows[1:]) if b - a >= 0.5 * lh]
    if gaps:
        _, a, b = max(gaps)
        return _xy_cut([w for w in items if w["bottom"] <= a], lh, depth + 1) + \
            _xy_cut([w for w in items if w["top"] >= b], lh, depth + 1)
    return [items]


def _pdf(path: str, out_dir: Path, reader: Reader | None, progress, key_only: bool = False) -> dict:
    import pdfplumber
    pages, tables, jobs = [], [], []
    with pdfplumber.open(path) as pdf:
        n_pages = len(pdf.pages)
        for pno, page in enumerate(pdf.pages, start=1):
            progress(0.05 + 0.25 * pno / n_pages, f"Reading page {pno} of {n_pages}")
            words = page.extract_words(extra_attrs=["size"], return_chars=True)  # font per character: a ligature's
            for w in words:                                                     # own font mustn't split a word
                chars = w.pop("chars", None) or []
                w["bold"] = bool(chars) and sum(bool(_BOLD.search(c.get("fontname", ""))) for c in chars) > len(chars) / 2
            words = _respace(words)
            lines = _lines(words)
            area = float(page.width * page.height)
            imgs = [(float(i["x0"]), float(i["top"]), float(i["x1"]), float(i["bottom"])) for i in page.images
                    if i["x1"] - i["x0"] >= MIN_IMAGE_PT[0] and i["bottom"] - i["top"] >= MIN_IMAGE_PT[1]]
            blocks = []  # (top, kind, payload)
            if len(words) < 5 and any((b[2] - b[0]) * (b[3] - b[1]) > 0.4 * area for b in imgs):
                t = {"id": f"p{pno:03d}-page", "page": pno, "source": "page image", "where": f"page {pno}",
                     "bbox": [0, 0, float(page.width), float(page.height)]}
                jobs.append((t, _png(page.to_image(resolution=DPI).original), None))
                tables.append(t)
                blocks.append((0, "table", t["id"]))
            else:
                ruled = [tuple(float(v) for v in tb.bbox) for tb in page.find_tables()]
                ruled = [b for b in ruled if (b[2] - b[0]) > 60 and (b[3] - b[1]) > 20]
                fixed = [(b, "text-layer table") for b in ruled]
                # borderless tables across the page first: they are solid blocks, never cut into columns
                fixed += [(b, "text-layer table") for b in _number_runs(_lines(words), ruled + imgs)]
                fixed += [(b, "picture") for b in imgs if not any(_overlap(b, r) for r, _ in fixed)]
                boxes = [b for b, _ in fixed]
                in_box = lambda w: any(b[0] - 1 <= (w["x0"] + w["x1"]) / 2 <= b[2] + 1 and b[1] - 1 <= (w["top"] + w["bottom"]) / 2 <= b[3] + 1
                                       for b in boxes)
                items = [w for w in words if not in_box(w)]
                items += [{"x0": b[0], "top": b[1], "x1": b[2], "bottom": b[3], "box": k} for k, b in enumerate(boxes)]
                heights = sorted(w["bottom"] - w["top"] for w in words) or [10]
                # reading order, then per block: its lines, the borderless tables among them, the fixed boxes
                ordered = []  # (kind, payload): ("line", line) / ("region", (box, source))
                for leaf in _xy_cut(items, heights[len(heights) // 2]):
                    lns = _lines([w for w in leaf if "box" not in w])
                    runs = _number_runs(lns, boxes)  # tables inside a column
                    entries = [(b[1], "region", (b, "text-layer table")) for b in runs]
                    entries += [(fixed[w["box"]][0][1], "region", fixed[w["box"]]) for w in leaf if "box" in w]
                    entries += [(ln["top"], "line", ln) for ln in lns if not any(_inside(ln, b) for b in runs)]
                    ordered += [(kind, payload) for _, kind, payload in sorted(entries, key=lambda e: e[0])]
                sizes = sorted(ln["size"] for kind, ln in ordered if kind == "line") or [10]
                body = sizes[len(sizes) // 2]
                prev, k = None, 0
                for seq, (kind, payload) in enumerate(ordered):
                    if kind == "region":
                        box, source = payload
                        k += 1
                        box = _pad(box, page)
                        t = {"id": f"p{pno:03d}-t{k}", "page": pno, "source": source, "where": f"page {pno}",
                             "bbox": [round(v, 1) for v in box]}
                        inside = [ln["text"] for ln in _lines([w for w in words if _inside(
                            {"top": w["top"], "bottom": w["bottom"], "x0": w["x0"], "x1": w["x1"]}, box)])]
                        if source == "text-layer table":
                            t["text_lines"] = inside  # kept so a person's edit can be checked against the page too
                        png = _png(page.crop(box).to_image(resolution=DPI).original)
                        jobs.append((t, png, inside if source == "text-layer table" and inside else None))
                        tables.append(t)
                        blocks.append((seq, "table", t["id"]))
                        prev = None
                        continue
                    ln = payload
                    if re.fullmatch(r"(page )?\d+( of \d+)?", ln["text"].strip(), re.I) and ln["top"] > page.height * 0.9:
                        continue  # page number in the footer
                    # a heading: bigger type, or a short bold line of its own (a side heading at body size)
                    heading = ln["size"] >= body * 1.3 or (ln["bold"] and len(ln["text"].split()) <= 8
                                                           and not ln["text"].rstrip().endswith((".", ",", ";")))
                    # a new paragraph after a wider gap, or where reading moves up to the next column
                    gap = prev is None or ln["top"] < prev["top"] - 1 or \
                        ln["top"] - prev["bottom"] > 1.2 * (prev["bottom"] - prev["top"])
                    blocks.append((seq, "heading" if heading else "text", ln["text"], gap))
                    prev = ln
            pages.append({"n": pno, "blocks": sorted(blocks, key=lambda b: b[0])})
    _read_all(reader, jobs, out_dir, progress, pages, key_only)
    return {"kind": "pdf", "pages": pages, "tables": tables}


def _overlap(a, b) -> bool:
    return not (a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1])


def _pad(box, page, pad: float = 6):
    return (max(0, box[0] - pad), max(0, box[1] - pad), min(float(page.width), box[2] + pad),
            min(float(page.height), box[3] + pad))


def _page_md(blocks, by_id) -> tuple[str, str | None]:
    out, title, para = [], None, []

    def end_para():
        if para:
            out.append(" ".join(para))
            para.clear()

    for b in blocks:
        if b[1] == "table":
            end_para()
            t = by_id[b[2]]
            if t.get("source") == "page image":
                title = title or t.get("title")
            out.append(table_md(t))
        elif b[1] == "heading":
            end_para()
            title = title or b[2]
            out.append(f"### {b[2]}")
        else:
            if b[3]:
                end_para()
            para.append(b[2])
    end_para()
    return "\n\n".join(x for x in out if x), title


def table_md(t: dict) -> str:
    """How a table appears in the document Markdown: the approved text if a person settled it."""
    if t.get("status") == "figure":
        return f"*[Figure: {t.get('description') or 'image'}]*"
    md = t.get("final_markdown") or t.get("markdown") or ""
    if not md and t.get("status") == "unread" and t.get("text_lines"):
        # not read yet: the PDF's own text for now, so the facts can already be found and quoted from it
        return f"<!-- table {t['id']} (unread) -->\n" + "  \n".join(t["text_lines"])
    if not md:
        return f"*[Table {t['id']}: not read yet]*" if t.get("deferred") else f"*[Table {t['id']}: not read]*"
    if t.get("source") == "page image":
        return md
    cap = f"**{t['title']}**\n\n" if t.get("title") else ""
    return f"<!-- table {t['id']} ({t.get('status')}) -->\n{cap}{md}"


# ---- which tables first -----------------------------------------------------------------------------------------
# The steps after the report need its key figures (target, valuation date, conclusions, key assumptions), which
# sit in a few tables. Those are read first; the report counts as read once they are, and the rest are read in
# the background (engagement.py) while the facts are extracted and the models are worked on.
KEY_WORDS = re.compile(r"valuation|enterprise value|equity value|discount rate|\bwacc\b|terminal|growth|multiple|"
                       r"net debt|preferred|\brange\b|sensitivit|conclu|assumption|cost of (?:equity|capital|debt)|"
                       r"gearing|\bbeta\b|risk premium|summary", re.I)
KEY_TABLES = 12


def _page_texts(pages: list[dict]) -> dict[int, str]:
    return {p["n"]: " ".join(str(b[2]) for b in p["blocks"] if len(b) > 2 and b[1] in ("text", "heading"))
            + " " + (p.get("title") or "") for p in pages}


def rank_tables(tables: list[dict], pages: list[dict]) -> list[tuple[int, dict]]:
    """[(score, table)], most likely to hold key figures first: valuation words in the table's own text count
    three times those on its page; then page order."""
    text = _page_texts(pages)
    words = lambda s: {m.lower() for m in KEY_WORDS.findall(s or "")}
    scored = [(3 * len(words(" ".join(t.get("text_lines") or []) + " " + (t.get("title") or "")))
               + len(words(text.get(t["page"], ""))), t) for t in tables]
    return sorted(scored, key=lambda x: (-x[0], x[1]["page"], x[1]["id"]))


def _read_all(reader: Reader | None, jobs: list, out_dir: Path, progress, pages: list | None = None,
              key_only: bool = False) -> None:
    (out_dir / "tables").mkdir(parents=True, exist_ok=True)
    for t, png, _ in jobs:
        t["png"] = f"tables/{t['id']}.png"
        (out_dir / t["png"]).write_bytes(png)
    if not jobs:
        return
    if reader is None:
        for t, _, _ in jobs:
            t.update(status="unread", markdown="", check=None)
        return
    if key_only and len(jobs) > KEY_TABLES:
        key = {t["id"] for score, t in rank_tables([j[0] for j in jobs], pages or [])[:KEY_TABLES] if score > 0}
        for t, _, _ in jobs:
            if t["id"] not in key:
                t.update(status="unread", markdown="", check=None, deferred=True)
        jobs = [j for j in jobs if j[0]["id"] in key]
    done = 0

    def one(job):
        t, png, lines = job
        try:
            t.update(read_and_check(reader, t, png, lines))
        except Exception as e:  # keep going; the table shows as failed with the reason
            t.update(status="error", markdown="", error=f"{type(e).__name__}: {e}")
        return t

    with ThreadPoolExecutor(READERS) as pool:
        for t in pool.map(calllog.carry(one), jobs):
            done += 1
            progress(0.3 + 0.65 * done / len(jobs), f"Read and checked {done} of {len(jobs)} tables")


# ---- PPTX ---------------------------------------------------------------------------------------------------

def _shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(s.shapes)
        else:
            yield s


def _grid_md(rows: list[list[str]]) -> str:
    rows = [[(c or "").replace("|", "/").replace("\n", " ").strip() for c in r] for r in rows]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    return "\n".join(["| " + " | ".join(rows[0]) + " |", "|" + "---|" * w] + ["| " + " | ".join(r) + " |" for r in rows[1:]])


def _chart_md(chart) -> str:
    plot = chart.plots[0]
    cats = [str(c) for c in plot.categories]
    series = [(s.name, list(s.values)) for p in chart.plots for s in p.series]
    rows = [[""] + [name for name, _ in series]]
    for i, c in enumerate(cats):
        rows.append([c] + [("" if i >= len(v) or v[i] is None else f"{v[i]:,.6g}") for _, v in series])
    return _grid_md(rows)


def _pptx(path: str, out_dir: Path, reader: Reader | None, progress, key_only: bool = False) -> dict:
    from PIL import Image
    from pptx import Presentation
    prs = Presentation(path)
    pages, tables, jobs = [], [], []
    n_slides = len(prs.slides)
    for sno, slide in enumerate(prs.slides, start=1):
        progress(0.05 + 0.25 * sno / n_slides, f"Reading slide {sno} of {n_slides}")
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else None
        blocks, k = [], 0
        for s in sorted(_shapes(slide.shapes), key=lambda s: ((s.top or 0), (s.left or 0))):
            top = s.top or 0
            if s is title_shape or (title_shape is not None and s.shape_id == title_shape.shape_id):
                continue
            if getattr(s, "has_table", False) and s.has_table:
                k += 1
                grid = [[c.text for c in r.cells] for r in s.table.rows]
                t = {"id": f"s{sno:03d}-t{k}", "page": sno, "source": "pptx table", "where": f"slide {sno}",
                     "is_table": True, "markdown": _grid_md(grid), "status": "verified",
                     "check": {"method": "native", "ok": True, "note": "read from the slide's table, not from an image"}}
                tables.append(t)
                blocks.append((top, "table", t["id"]))
            elif getattr(s, "has_chart", False) and s.has_chart:
                k += 1
                ch = s.chart
                name = ch.chart_title.text_frame.text if ch.has_title and ch.chart_title.has_text_frame else None
                t = {"id": f"s{sno:03d}-t{k}", "page": sno, "source": "pptx chart data", "where": f"slide {sno}",
                     "is_table": True, "title": name, "markdown": _chart_md(ch), "status": "verified",
                     "check": {"method": "native", "ok": True, "note": "the chart's own data, from the file"}}
                tables.append(t)
                blocks.append((top, "table", t["id"]))
            elif hasattr(s, "image"):
                try:
                    img = Image.open(io.BytesIO(s.image.blob))
                except Exception:
                    continue
                if img.width < MIN_PICTURE_PX[0] or img.height < MIN_PICTURE_PX[1]:
                    continue
                k += 1
                t = {"id": f"s{sno:03d}-t{k}", "page": sno, "source": "picture", "where": f"slide {sno}"}
                jobs.append((t, _png(img), None))
                tables.append(t)
                blocks.append((top, "table", t["id"]))
            elif s.has_text_frame and s.text_frame.text.strip():
                lines = []
                for p in s.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs).strip()
                    if txt:
                        lines.append(("  " * max(0, p.level - 1) + "- " if p.level else "") + txt)
                blocks.append((top, "text", "\n".join(lines), True))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            blocks.append((10 ** 12, "text", "> Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip(), True))
        pages.append({"n": sno, "title": title, "blocks": sorted(blocks, key=lambda b: b[0])})
    _read_all(reader, jobs, out_dir, progress, pages, key_only)
    return {"kind": "pptx", "pages": pages, "tables": tables}


# ---- entry point --------------------------------------------------------------------------------------------

def render(doc: dict) -> str:
    """(Re)build each page's Markdown and the whole document's from the page blocks and the tables, so a
    table a person approves or edits shows up everywhere. Slides count as pages; <!-- page N --> marks each."""
    by_id = {t["id"]: t for t in doc["tables"]}
    label = "Slide" if doc["kind"] == "pptx" else "Page"
    for p in doc["pages"]:
        md, title = _page_md(p["blocks"], by_id)
        if doc["kind"] == "pptx":
            md = (f"### {p['title']}\n\n" if p.get("title") else "") + md
        else:
            p["title"] = title
        p["md"] = md
    doc["markdown"] = "\n\n".join(f"<!-- page {p['n']} -->\n## {label} {p['n']}\n\n{p['md']}" for p in doc["pages"])
    return doc["markdown"]


def read_rest(doc: dict, out_dir: str | Path, model: str, reviewer_model: str, on_usage=None, progress=None,
              on_table=None) -> int:
    """Read the tables process(key_only=True) left for later, most likely to matter first. on_table(t) is called
    as each one is done (so it can be saved at once). Returns how many were read."""
    progress = progress or (lambda frac, msg: None)
    todo = [t for score, t in rank_tables([t for t in doc["tables"] if t.get("deferred") and t.get("status") == "unread"
                                           and t.get("png")], doc["pages"])]
    if not todo:
        return 0
    reader = Reader(model, reviewer_model, on_usage)
    done = 0

    def one(t):
        try:
            png = (Path(out_dir) / t["png"]).read_bytes()
            lines = t.get("text_lines") if t["source"] == "text-layer table" else None
            t.update(read_and_check(reader, t, png, lines))
        except Exception as e:
            t.update(status="error", markdown="", error=f"{type(e).__name__}: {e}")
        t.pop("deferred", None)
        return t

    with ThreadPoolExecutor(READERS) as pool:
        for t in pool.map(calllog.carry(one), todo):
            done += 1
            progress(done / len(todo), f"Read and checked {done} of {len(todo)} remaining tables")
            if on_table:
                on_table(t)
    return done


def process(path: str, out_dir: str | Path, model: str = "gpt-6-luna", reviewer_model: str = "gpt-6-sol",
            progress=None, on_usage=None, read: bool = True, key_only: bool = False) -> dict:
    """Read a report. key_only: read only the tables most likely to hold the key figures now (up to KEY_TABLES);
    the others are marked deferred for read_rest()."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    progress = progress or (lambda frac, msg: None)
    reader = Reader(model, reviewer_model, on_usage) if read else None
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        doc = _pdf(path, out_dir, reader, progress, key_only)
    elif ext == ".pptx":
        doc = _pptx(path, out_dir, reader, progress, key_only)
    else:
        raise ValueError(f"{Path(path).name}: reports must be .pdf or .pptx (save .ppt / .docx as PDF first)")
    (out_dir / "document.md").write_text(render(doc), encoding="utf-8")
    progress(1.0, "Document read")
    return doc


if __name__ == "__main__":
    sys.path.insert(0, str(Path(__file__).parent))
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    d = process(sys.argv[1], sys.argv[2], *(sys.argv[3:4] * 2), progress=lambda f, m: print(f"{f:4.0%} {m}"))
    for t in d["tables"]:
        print(t["id"], t["source"], t["status"], json.dumps(t.get("check") or t.get("second_read"), default=str)[:300])
    print(d["markdown"][:3000])
