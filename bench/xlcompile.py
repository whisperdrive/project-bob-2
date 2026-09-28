"""Compile an overlay's Excel formulas into a readable Python module: one function per line item.

Every formula cell on the chosen sheets becomes Python. Cells in a row whose formulas are the same apart from
a column shift (the usual timeline row: =D5+D6, =E5+E6, ...) share one branch of that row's function, which
takes the column as c:

    @B.row('DCF', 7, [(4, 23)], 'Valuation cash flow [A$m]')
    def DCF_7(c):
        if 4 <= c <= 23:  # D:W  =D5+D6
            return xl.add(DCF_5(c), DCF_6(c))
        return B.input('DCF', 7, c)

References to other overlay rows call their functions; constants are the overlay's inputs (B.input); cells
outside the overlay come from the client model through B.feed_value (same workbook) or B.ext_value (an
external link [n]). IF / IFERROR / IFNA / CHOOSE branches are lambdas, so only the branch taken is computed.
Defined names, whole-row / whole-column references and ISFORMULA are resolved here. Workbook text only enters
the code through repr(), so a label or string can't become code.
    uv run python bench/xlcompile.py out/<dir>/model.db Sheet1 [Sheet2 ...] > overlay.py
"""
import re
import sqlite3
import sys
from collections import defaultdict

from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string, get_column_letter

import rodb
import xlruntime

KNOWN = {n for n in dir(xlruntime.xl) if n.isupper() and not n.startswith("_")} - {"ERR", "MISSING", "INDEX_REF", "ISECT", "RANGE"}
OPS = {"+": "add", "-": "sub", "*": "mul", "/": "div", "^": "pow", "&": "concat", "=": "eq", "<>": "ne", "<": "lt",
       ">": "gt", "<=": "le", ">=": "ge"}
BP = {":": 90, " ": 80, "^": 50, "*": 40, "/": 40, "+": 30, "-": 30, "&": 20,
      "=": 10, "<>": 10, "<": 10, ">": 10, "<=": 10, ">=": 10}
# Arguments that are references, not values: passed to the function as ranges (Rng), evaluated lazily.
ALL = "all"
REF_ARGS = {"OFFSET": {0}, "INDEX": {0}, "ROW": {0}, "COLUMN": {0}, "ROWS": {0}, "COLUMNS": {0}, "CELL": {1},
            "MATCH": {1}, "VLOOKUP": {1}, "HLOOKUP": {1}, "LOOKUP": {1, 2}, "XNPV": {1, 2}, "XIRR": {0, 1},
            "IRR": {0}, "N": {0}, "TRANSPOSE": {0}, "COUNTBLANK": {0}}
for _f in ("SUMIFS", "SUMIF", "COUNTIFS", "COUNTIF", "AVERAGEIFS", "AVERAGEIF", "MAXIFS", "MINIFS", "SUMPRODUCT",
           "MMULT", "SUM", "MAX", "MIN", "AVERAGE", "COUNT", "COUNTA", "PRODUCT", "SUMSQ", "AND", "OR", "XOR",
           "CONCAT", "NPV"):
    REF_ARGS[_f] = ALL
VOLATILE = {"TODAY", "NOW", "RAND", "RANDBETWEEN", "INDIRECT", "CELL", "INFO"}
MAX_DEPTH = 150  # Python's parser refuses deeper nesting; such a formula is reported, not compiled


class CompileError(Exception):
    pass


# ---- tokens and parsing -------------------------------------------------------------------------------------

def tokens(formula: str) -> list[tuple[str, str]]:
    out = []
    for t in Tokenizer(formula).items:
        v = t.value
        if t.type == Token.FUNC:
            if t.subtype == Token.OPEN:
                # "A5:OFFSET(" or ":INDEX(": a range operator in front of a function
                m = re.match(r"^(.*):([^:]+\()$", v)
                if m:
                    if m.group(1):
                        out.append(("range", m.group(1)))
                    out.append(("op", ":"))
                    v = m.group(2)
                out.append(("func", v[:-1]))
            else:
                out.append(("close", ")"))
        elif t.type == Token.PAREN:
            out.append(("lparen" if t.subtype == Token.OPEN else "rparen", v))
        elif t.type == Token.SEP:
            out.append(("sep" if t.subtype == Token.ARG else "rowsep", v))
        elif t.type == Token.ARRAY:
            out.append(("lbrace" if t.subtype == Token.OPEN else "rbrace", v))
        elif t.type == Token.OP_PRE:
            out.append(("prefix", v))
        elif t.type == Token.OP_IN:
            out.append(("op", v))
        elif t.type == Token.OP_POST:
            out.append(("postfix", v))
        elif t.type == Token.WSPACE:
            out.append(("ws", v))
        elif t.type == Token.OPERAND:
            if t.subtype == Token.RANGE and v.startswith(":"):
                out.append(("op", ":"))
                v = v[1:]
            out.append((t.subtype.lower(), v))
    res = []
    for i, tk in enumerate(out):  # a space between two references is the intersection operator
        if tk[0] == "ws":
            nxt = out[i + 1] if i + 1 < len(out) else None
            if res and nxt and res[-1][0] in ("range", "close", "rparen") and nxt[0] in ("range", "func", "lparen"):
                res.append(("op", " "))
            continue
        res.append(tk)
    return res


class Parser:
    def __init__(self, toks):
        self.t, self.i = toks, 0

    def peek(self):
        return self.t[self.i] if self.i < len(self.t) else ("end", None)

    def next(self):
        tk = self.peek()
        self.i += 1
        return tk

    def parse(self):
        node = self.expr(0)
        if self.peek()[0] != "end":
            raise CompileError(f"unexpected {self.peek()[1]!r}")
        return node

    def expr(self, rbp):
        left = self.prefix()
        while True:
            kind, v = self.peek()
            if kind == "postfix":
                if rbp >= 70:
                    break
                self.next()
                left = ("pct", left)
                continue
            if kind != "op" or BP.get(v, 0) <= rbp:
                break
            self.next()
            left = ("op", v, left, self.expr(BP[v]))
        return left

    def prefix(self):
        kind, v = self.next()
        if kind == "prefix":
            inner = self.expr(55)  # Excel: negation binds tighter than ^, so -2^2 = 4
            return ("neg", inner) if v == "-" else inner
        if kind == "number":
            return ("n", float(v))
        if kind == "text":
            return ("s", v[1:-1].replace('""', '"'))
        if kind == "logical":
            return ("b", v.upper() == "TRUE")
        if kind == "error":
            return ("e", v.upper())
        if kind == "range":
            return ("ref", v)
        if kind == "lparen":
            e = self.expr(0)
            if self.next()[0] != "rparen":
                raise CompileError("unions like (A1,B1) aren't supported")
            return e
        if kind == "func":
            name = v.upper()
            for p in ("_XLFN._XLWS.", "_XLFN.", "_XLWS."):
                name = name.removeprefix(p)
            return ("f", name, self.args())
        if kind == "lbrace":
            return ("arr", self.array())
        raise CompileError(f"unexpected {v!r}")

    def args(self):
        args = []
        if self.peek()[0] == "close":
            self.next()
            return args
        while True:
            args.append(("missing",) if self.peek()[0] in ("sep", "close") else self.expr(0))
            kind, _ = self.next()
            if kind == "close":
                return args
            if kind != "sep":
                raise CompileError("bad argument list")

    def array(self):
        rows, row = [], []
        while True:
            row.append(self.expr(0))
            kind, _ = self.next()
            if kind == "sep":
                continue
            rows.append(row)
            if kind == "rowsep":
                row = []
                continue
            if kind == "rbrace":
                return rows
            raise CompileError("bad array constant")


def parse(formula: str):
    return Parser(tokens(formula)).parse()


# ---- references ---------------------------------------------------------------------------------------------

_SHEET = re.compile(r"^(?:'((?:[^']|'')+)'|([^'!]+))!(.+)$")
_CELL = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_COLS = re.compile(r"^(\$?)([A-Za-z]{1,3}):(\$?)([A-Za-z]{1,3})$")
_ROWS = re.compile(r"^(\$?)(\d+):(\$?)(\d+)$")


def parse_ref(text: str) -> dict | None:
    """'[1]CashFlow'!$D$9:W9 -> {ext, sheet, kind, r1, c1, r2, c2, cabs1, cabs2}; None if it's a defined name."""
    ext, sheet, addr = None, None, text
    m = _SHEET.match(text)
    if m:
        sheet = (m.group(1) or "").replace("''", "'") if m.group(1) is not None else m.group(2)
        addr = m.group(3)
        em = re.match(r"^\[(\d+)\](.*)$", sheet)
        if em:
            ext, sheet = int(em.group(1)), em.group(2)
    if "#REF!" in addr.upper():
        return {"error": "#REF!"}
    parts = addr.split(":")
    if len(parts) == 1:
        cm = _CELL.match(addr)
        if not cm or column_index_from_string(cm.group(2).upper()) > 16384:
            return None if sheet is None else {"name": addr, "sheet": sheet, "ext": ext}
        c, r = column_index_from_string(cm.group(2).upper()), int(cm.group(4))
        return {"ext": ext, "sheet": sheet, "kind": "cell", "r1": r, "r2": r, "c1": c, "c2": c,
                "cabs1": bool(cm.group(1)), "cabs2": bool(cm.group(1))}
    if len(parts) != 2:
        return None
    a, b = _CELL.match(parts[0]), _CELL.match(parts[1])
    if a and b:
        ca, cb = column_index_from_string(a.group(2).upper()), column_index_from_string(b.group(2).upper())
        return {"ext": ext, "sheet": sheet, "kind": "range", "r1": int(a.group(4)), "r2": int(b.group(4)), "c1": ca,
                "c2": cb, "cabs1": bool(a.group(1)), "cabs2": bool(b.group(1))}
    m = _COLS.match(addr)
    if m:
        return {"ext": ext, "sheet": sheet, "kind": "cols", "r1": None, "r2": None,
                "c1": column_index_from_string(m.group(2).upper()), "c2": column_index_from_string(m.group(4).upper()),
                "cabs1": bool(m.group(1)), "cabs2": bool(m.group(3))}
    m = _ROWS.match(addr)
    if m:
        return {"ext": ext, "sheet": sheet, "kind": "rows", "r1": int(m.group(2)), "r2": int(m.group(4)), "c1": None,
                "c2": None, "cabs1": True, "cabs2": True}
    return None


def group_key(formula: str, col: int) -> str:
    """Formula text with relative columns written as offsets from the cell's column: cells in a row with the same
    key compute the same thing shifted along, so they share code."""
    out = []
    try:
        toks = tokens(formula)
    except Exception:
        return f"!{col}!{formula}"
    for kind, v in toks:
        if kind == "range":
            r = parse_ref(v)
            if r and "kind" in r and r["kind"] != "rows":
                c1 = r["c1"] if r["cabs1"] else f"~{r['c1'] - col}"
                c2 = r["c2"] if r["cabs2"] else f"~{r['c2'] - col}"
                v = f"{r['ext']}|{r['sheet']}|{r['kind']}|{r['r1']}|{c1}|{r['r2']}|{c2}"
        out.append(f"{kind}:{v}")
    return "\x1f".join(out)


# ---- code generation ----------------------------------------------------------------------------------------

class Gen:
    """AST -> Python expression for the cell(s) at (sheet, row, column c)."""

    def __init__(self, ctx: dict):
        self.ctx = ctx
        self.isformula = ctx["isformula_targets"]
        self.issues = []

    def colexpr(self, target: int, absolute: bool) -> str:
        if absolute:
            return str(target)
        off = target - self.col
        return "c" if off == 0 else (f"c + {off}" if off > 0 else f"c - {-off}")

    def ref(self, r: dict, want: str) -> str:
        if "error" in r:
            return "xl.ERR['#REF!']"
        sheet = r["sheet"] if r["sheet"] is not None else self.sheet
        src = r["ext"] if r["ext"] is not None else ""
        if src == "" and sheet not in self.ctx["sheets"]:
            return "xl.ERR['#REF!']"
        maxr, maxc = self.ctx["extent"].get((src, sheet), (1, 1))
        kind = r["kind"]
        r1, r2 = (1, maxr) if kind == "cols" else (r["r1"], r["r2"])
        if kind == "rows":
            c1, c2 = "1", str(maxc)
        else:
            c1, c2 = self.colexpr(r["c1"], r["cabs1"]), self.colexpr(r["c2"], r["cabs2"])
        if kind == "cell" and want == "value":
            if src == "":
                if sheet in self.ctx["overlay"]:
                    fn = self.ctx["fn"].get((sheet, r1))
                    if not fn:  # an input: name it in the branch's comment
                        where = f"{sheet}!{get_column_letter(r['c1'])}{r1}" if r["cabs1"] else f"{sheet}!r{r1}"
                        self.inputs_used.setdefault(where, self.ctx["labels"].get((sheet, r1), ("", ""))[0])
                    return f"{fn}({c1})" if fn else f"B.input({sheet!r}, {r1}, {c1})"
                return f"B.feed_value({sheet!r}, {r1}, {c1})"
            return f"B.ext_value({src}, {sheet!r}, {r1}, {c1})"
        return f"B.rng({src!r}, {sheet!r}, {r1}, {c1}, {r2}, {c2})"

    def name(self, name: str, want: str, sheet: str | None = None) -> str:
        """A defined name: the sheet's own (the formula's sheet, or the one it names) before the workbook's."""
        scope = sheet or self.sheet
        key = (scope, name.lower())
        target = self.ctx["local_names"].get(key)
        if target is None:
            target, key = self.ctx["names"].get(name.lower()), (None, name.lower())
        if target is None or key in self.names_seen:
            self.ctx["missing_names"].add(f"{scope}!{name}" if sheet else name)
            return "xl.ERR['#NAME?']"
        self.names_seen.add(key)
        try:
            return self.expr(parse("=" + target.lstrip("=")), want)
        finally:
            self.names_seen.discard(key)

    def lazy(self, node, want="value") -> str:
        return f"lambda: {self.expr(node, want)}"

    def fn(self, name: str, args: list, want: str) -> str:
        if name == "IF":
            if not args:
                raise CompileError("IF without arguments")
            yes = self.lazy(args[1], want) if len(args) > 1 else "None"
            no = f", {self.lazy(args[2], want)}" if len(args) > 2 else ""
            return f"xl.IF({self.expr(args[0])}, {yes}{no})"
        if name in ("IFERROR", "IFNA"):
            return f"xl.{name}({self.lazy(args[0], want)}, {self.lazy(args[1], want)})"
        if name == "CHOOSE":
            return f"xl.CHOOSE({self.expr(args[0])}, " + ", ".join(self.lazy(a, want) for a in args[1:]) + ")"
        if name == "ISFORMULA":
            a = args[0] if args else None
            r = parse_ref(a[1]) if a and a[0] == "ref" else None
            if not r or r.get("kind") != "cell" or r.get("ext") is not None:
                self.issues.append("ISFORMULA of something other than one cell in this workbook")
                return "xl.ERR['#VALUE!']"
            sheet = r["sheet"] or self.sheet
            for col in self.cols:  # every column this code serves
                target = r["c1"] if r["cabs1"] else r["c1"] + col - self.col
                self.isformula.add((sheet, r["r1"], target))
            return f"(({sheet!r}, {r['r1']}, {self.colexpr(r['c1'], r['cabs1'])}) in ISFORMULA)"
        if name == "IFS":  # condition, value pairs, in order: a chain of lazy IFs, #N/A if none holds
            if len(args) < 2 or len(args) % 2:
                raise CompileError("IFS needs condition, value pairs")
            code = "xl.ERR['#N/A']"
            for cond, val in reversed(list(zip(args[::2], args[1::2]))):
                code = f"xl.IF({self.expr(cond)}, {self.lazy(val, want)}, lambda: {code})"
            return code
        if name == "SWITCH":  # expression, then value, result pairs, then an optional default
            if len(args) < 3:
                raise CompileError("SWITCH needs an expression and a value, result pair")
            e, rest = self.expr(args[0]), args[1:]
            code = self.expr(rest[-1], want) if len(rest) % 2 else "xl.ERR['#N/A']"
            pairs = list(zip(rest[::2], rest[1::2]))
            for val, res in reversed(pairs):
                code = f"xl.IF(xl.eq({e}, {self.expr(val)}), {self.lazy(res, want)}, lambda: {code})"
            return code
        if name in ("ROW", "COLUMN") and not args:
            return f"{float(self.row)!r}" if name == "ROW" else "float(c)"
        if name in VOLATILE:
            self.ctx["volatile"].add(name)
        spec = REF_ARGS.get(name, set())
        parts = []
        for i, a in enumerate(args):
            if a[0] == "missing":
                parts.append("xl.MISSING")
            else:
                parts.append(self.expr(a, "ref" if spec == ALL or i in spec else "value"))
        if name == "INDEX" and want == "ref":
            return f"xl.INDEX_REF({', '.join(parts)})"
        py = name.replace(".", "_")  # STDEV.S -> STDEV_S
        if py in KNOWN:
            return f"xl.{py}({', '.join(parts)})"
        self.ctx["unknown"].add(name)
        return f"B.unknown({name!r})({', '.join(parts)})"

    def expr(self, node, want="value") -> str:
        k = node[0]
        if k == "n":
            return repr(node[1])
        if k == "s":
            return repr(node[1])
        if k == "b":
            return "True" if node[1] else "False"
        if k == "e":
            return f"xl.ERR[{node[1]!r}]" if node[1] in xlruntime.ERR else "xl.ERR['#VALUE!']"
        if k == "missing":
            return "xl.MISSING"
        if k == "ref":
            r = parse_ref(node[1])
            if r is None:
                return self.name(node[1], want)
            if "name" in r:
                if r.get("ext") is None and r.get("sheet"):  # Sheet!Name: that sheet's own name
                    return self.name(r["name"], want, r["sheet"])
                self.issues.append(f"external name {node[1]}")
                self.ctx["missing_names"].add(node[1])
                return "xl.ERR['#NAME?']"
            return self.ref(r, want)
        if k == "f":
            return self.fn(node[1], node[2], want)
        if k == "op":
            op, a, b = node[1], node[2], node[3]
            if op == ":":
                return f"xl.RANGE({self.expr(a, 'ref')}, {self.expr(b, 'ref')})"
            if op == " ":
                return f"xl.ISECT({self.expr(a, 'ref')}, {self.expr(b, 'ref')})"
            return f"xl.{OPS[op]}({self.expr(a)}, {self.expr(b)})"
        if k == "neg":
            return f"xl.neg({self.expr(node[1])})"
        if k == "pct":
            return f"xl.pct({self.expr(node[1])})"
        if k == "arr":
            return "xl.array([" + ", ".join("[" + ", ".join(self.expr(x) for x in row) + "]" for row in node[1]) + "])"
        raise CompileError(f"unknown node {k}")

    def cell(self, formula: str, sheet: str, row: int, col: int, cols: list[int]) -> str:
        self.sheet, self.row, self.col, self.cols = sheet, row, col, cols
        self.names_seen = set()
        self.inputs_used = {}
        return self.expr(parse(formula))


def _depth(code: str) -> int:
    d = m = 0
    for ch in code:
        if ch in "([{":
            d += 1
            m = max(m, d)
        elif ch in ")]}":
            d -= 1
    return m


def _ident(sheet: str, taken: set) -> str:
    s = re.sub(r"\W+", "_", sheet).strip("_") or "Sheet"
    if s[0].isdigit():
        s = "S_" + s
    base, n = s, 2
    while s in taken:
        s, n = f"{base}_{n}", n + 1
    taken.add(s)
    return s


def _runs(cols: list[int]) -> list[tuple[int, int]]:
    runs = []
    for c in sorted(cols):
        if runs and c == runs[-1][1] + 1:
            runs[-1][1] = c
        else:
            runs.append([c, c])
    return [tuple(r) for r in runs]


def _cond(runs) -> str:
    parts = [f"c == {a}" if a == b else f"{a} <= c <= {b}" for a, b in runs]
    return " or ".join(parts)


def _comment(text: str, n: int = 110) -> str:
    text = re.sub(r"[\r\n\t]+", " ", text)
    return text if len(text) <= n else text[: n - 1] + "…"


def compile_overlay(db_path: str, sheets: list[str], title: str = "", progress=None,
                    source_path: str | None = None) -> tuple[str, dict]:
    """(module source, stats) for the formula cells on `sheets` of the workbook behind model.db. source_path: the
    workbook itself, to read its sheet-level defined names when model.db was built before it kept them."""
    progress = progress or (lambda f, m: None)
    db = rodb.connect(db_path)
    wb_sheets = [s for (s,) in db.execute("SELECT sheet FROM sheets ORDER BY rowid")]
    overlay = [s for s in wb_sheets if s in set(sheets)]
    have = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    extent = {("", s): (r or 1, c or 1) for s, r, c in db.execute("SELECT sheet, MAX(row), MAX(col) FROM cells GROUP BY sheet")}
    if "extcells" in have:
        extent.update({(i, s): (r or 1, c or 1) for i, s, r, c in
                       db.execute("SELECT idx, sheet, MAX(row), MAX(col) FROM extcells GROUP BY idx, sheet")})
    ncols = {r[1] for r in db.execute("PRAGMA table_info(names)")} if "names" in have else set()
    rows = db.execute("SELECT name, ref, scope FROM names").fetchall() if "scope" in ncols else \
        [(n, r, None) for n, r in db.execute("SELECT name, ref FROM names")] if "names" in have else []
    if "scope" not in ncols and source_path:
        try:
            import build_map
            rows = build_map.defined_names(source_path)
        except Exception:
            pass
    names = {n.lower(): ref for n, ref, sc in rows if not sc}
    local_names = {(sc, n.lower()): ref for n, ref, sc in rows if sc}
    labels = {(s, r): (lab or "", u or "") for s, r, lab, u in db.execute("SELECT sheet, row, label, units FROM rows")}

    cells = defaultdict(list)
    q = f"SELECT sheet, row, col, formula FROM cells WHERE formula IS NOT NULL AND sheet IN ({','.join('?' * len(overlay))})"
    for s, r, c, f in db.execute(q, overlay):
        cells[(s, r)].append((c, f))
    taken = set()
    ident = {s: _ident(s, taken) for s in overlay}
    fn = {(s, r): f"{ident[s]}_{r}" for (s, r) in cells}
    ctx = {"sheets": set(wb_sheets), "overlay": set(overlay), "extent": extent, "names": names, "fn": fn, "labels": labels,
           "isformula_targets": set(), "unknown": set(), "volatile": set(), "local_names": local_names,
           "missing_names": set()}
    gen = Gen(ctx)
    stats = {"formula_cells": 0, "rows": len(cells), "branches": 0, "not_compiled": []}
    body = []
    order = sorted(cells, key=lambda k: (overlay.index(k[0]), k[1]))
    for n, key in enumerate(order):
        if n % 200 == 0:
            progress(n / max(1, len(order)), f"Compiling {key[0]} row {key[1]}")
        s, r = key
        groups = defaultdict(list)
        first_formula = {}
        for c, f in sorted(cells[key]):
            k = group_key(f, c)
            groups[k].append(c)
            first_formula.setdefault(k, f)
        lab, units = labels.get(key, ("", ""))
        runs_all = _runs([c for cs in groups.values() for c in cs])
        label = f"{lab} [{units}]" if units else lab
        lines = [f"# {s}!r{r}  {_comment(label, 90)}" if label else f"# {s}!r{r}",
                 f"@B.row({s!r}, {r}, {runs_all!r}, {label!r})", f"def {fn[key]}(c):"]
        for k, cs in sorted(groups.items(), key=lambda kv: min(kv[1])):
            c0 = min(cs)
            stats["formula_cells"] += len(cs)
            stats["branches"] += 1
            runs = _runs(cs)
            where = ",".join(get_column_letter(a) if a == b else f"{get_column_letter(a)}:{get_column_letter(b)}"
                             for a, b in runs[:4]) + ("…" if len(runs) > 4 else "")
            try:
                code = gen.cell(first_formula[k], s, r, c0, cs)
                if _depth(code) > MAX_DEPTH:
                    raise CompileError("nested too deeply for Python")
            except Exception as e:
                code = "xl.ERR['#NAME?']"
                stats["not_compiled"].append({"cell": f"{s}!{get_column_letter(c0)}{r}", "formula": first_formula[k][:200],
                                              "reason": f"{type(e).__name__}: {e}"})
            lines.append(f"    if {_cond(runs)}:  # {where}  {_comment(first_formula[k])}")
            if gen.inputs_used:
                used = "; ".join(f"{k2} {_comment(lab, 40)}".strip() for k2, lab in list(gen.inputs_used.items())[:4])
                lines.append(f"        # inputs: {used}{' …' if len(gen.inputs_used) > 4 else ''}")
            lines.append(f"        return {code}")
        lines.append(f"    return B.input({s!r}, {r}, c)")
        body.append("\n".join(lines))
    progress(0.95, "Resolving ISFORMULA")
    isf = set()
    for (s, r, c) in sorted(ctx["isformula_targets"]):
        hit = db.execute("SELECT 1 FROM cells WHERE sheet=? AND row=? AND col=? AND formula IS NOT NULL", (s, r, c)).fetchone()
        if hit:
            isf.add((s, r, c))
    db.close()
    stats.update(unknown_functions=sorted(ctx["unknown"]), missing_names=sorted(ctx["missing_names"])[:100],
                 volatile=sorted(ctx["volatile"]), sheets=overlay,
                 issues=sorted(set(gen.issues)))
    head = f'''"""Python overlay{": " + _comment(title, 80) if title else ""}.

Compiled from the workbook's formulas by bench/xlcompile.py; sheets: {", ".join(overlay)}.
{stats["formula_cells"]:,} formula cells in {stats["rows"]:,} line items, {stats["branches"]:,} branches.
Each function is one line item (row): it returns the cell's value for column c. Constants are the overlay's
inputs (B.input); values from outside the overlay come from the client model (B.feed_value, B.ext_value).
Excel functions are in xl (bench/xlruntime.py). Regenerate rather than edit: the workbook is the source.
"""
from xlruntime import Book, xl

B = Book()
B.overlay = set({overlay!r})
ISFORMULA = {isf!r}
'''
    return head + "\n\n" + "\n\n\n".join(body) + "\n", stats


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    src, st = compile_overlay(sys.argv[1], sys.argv[2:])
    sys.stdout.write(src)
    print({k: v for k, v in st.items() if k != "not_compiled"}, len(st["not_compiled"]), file=sys.stderr)
